from pathlib import Path

import pytest
from docx import Document

from apple_fill.documents import LoadedDocument, Page
from apple_fill.extract import ExtractError, extract_product
from apple_fill.llm import model_output_to_product
from apple_fill.product import ProductError, require_product
from apple_fill.sample_data import PAGE_TWO, write_docx


def test_docx_extracts_literal_fields_and_leaves_unknowns_empty(tmp_path: Path):
    path = tmp_path / "spec.docx"
    write_docx(path)
    product = extract_product([path], "qp-200", backend="labels")
    name = product["fields"]["productName"]
    assert name["status"] == "sourced"
    assert name["value"] == "青苹空气净化器"
    assert name["sources"][0]["page"] == 1
    assert product["fields"]["description"]["value"] == PAGE_TWO.split("：", 1)[1]
    assert product["fields"]["description"]["sources"][0]["page"] == 2
    assert product["fields"]["productNameEn"]["status"] == "missing"
    assert product["fields"]["productNameEn"]["value"] is None
    assert product["fields"]["category"]["status"] == "missing"
    voltage = next(item for item in product["fields"]["specs"] if item["key"] == "ratedVoltage")
    assert voltage["value"] == "220V"


def test_table_rows_are_read(tmp_path: Path):
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "产品名称"
    table.rows[0].cells[1].text = "桌上风扇"
    document.save(path)
    product = extract_product([path], "fan", backend="labels")
    assert product["fields"]["productName"]["value"] == "桌上风扇"


def test_conflicting_documents_are_not_collapsed(tmp_path: Path):
    first = tmp_path / "a.docx"
    second = tmp_path / "b.docx"
    write_docx(first)
    document = Document()
    document.add_paragraph("型号：QP-300")
    document.save(second)
    product = extract_product([first, second], "qp", backend="labels")
    model = product["fields"]["modelNo"]
    assert model["status"] == "conflict"
    assert model["value"] is None
    assert {item["value"] for item in model["candidates"]} == {"QP-200", "QP-300"}


def test_doctored_value_fails_the_quote_check(tmp_path: Path):
    path = tmp_path / "spec.docx"
    write_docx(path)
    product = extract_product([path], "qp-200", backend="labels")
    product["fields"]["specs"][0]["value"] = "110V"
    from apple_fill.documents import read_pages

    pages = read_pages(path, "d1")
    with pytest.raises(ProductError):
        require_product(product, {"d1": {page.page: page.text for page in pages}})


def _document_with(text: str) -> LoadedDocument:
    return LoadedDocument(
        doc_id="d1",
        path=Path("spec.docx"),
        filename="spec.docx",
        sha256="a" * 64,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        pages=[Page("d1", 1, text)],
    )


def test_model_cannot_mark_a_literal_field_as_derived():
    document = _document_with("型号：QP-200")
    output = {"fields": {"modelNo": {"value": "QP-200", "status": "derived", "sources": []}}}
    with pytest.raises(ExtractError):
        model_output_to_product(output, [document], product_id="qp", model="test")


def test_model_sourced_value_must_match_the_quote():
    document = _document_with("额定电压：220V")
    output = {
        "fields": {
            "specs": [
                {
                    "key": "ratedVoltage",
                    "label": "额定电压",
                    "value": "110V",
                    "status": "sourced",
                    "sources": [{"docId": "d1", "page": 1, "quote": "额定电压：220V"}],
                }
            ]
        }
    }
    with pytest.raises(ExtractError, match="不能从引用里得到"):
        model_output_to_product(output, [document], product_id="qp", model="test")
