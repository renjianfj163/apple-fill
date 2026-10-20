"""Read PDF, Word, and JPEG into pages.

Page text comes from the file itself, or from Tesseract for images. Model
output is never copied into the page text it will later be checked against.
"""

from __future__ import annotations

import hashlib
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


class DocumentError(ValueError):
    pass


MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


@dataclass
class Page:
    doc_id: str
    page: int
    text: str


@dataclass
class LoadedDocument:
    doc_id: str
    path: Path
    filename: str
    sha256: str
    mime: str
    pages: list[Page] = field(default_factory=list)

    @property
    def text_available(self) -> bool:
        return any(page.text.strip() for page in self.pages)

    def as_dict(self) -> dict[str, object]:
        return {
            "docId": self.doc_id,
            "filename": self.filename,
            "sha256": self.sha256,
            "mime": self.mime,
            "pageCount": max(1, len(self.pages)),
            "textAvailable": self.text_available,
            "path": str(self.path),
        }

    def text_map(self) -> dict[int, str]:
        return {page.page: page.text for page in self.pages}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _paragraph_has_page_break(paragraph) -> bool:
    breaks = paragraph._p.findall(".//" + qn("w:br"))
    return any(item.get(qn("w:type")) == "page" for item in breaks)


def read_docx(path: Path, doc_id: str) -> list[Page]:
    document = Document(str(path))
    pages: list[list[str]] = [[]]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            pages[-1].append(text)
        if _paragraph_has_page_break(paragraph):
            pages.append([])
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                cell_text = " ".join(cell.text.split())
                if cell_text and cell_text not in cells:
                    cells.append(cell_text)
            if len(cells) >= 2:
                pages[0].append(f"{cells[0]}：{cells[1]}")
            elif cells:
                pages[0].append(cells[0])
    return [
        Page(doc_id, index, "\n".join(lines))
        for index, lines in enumerate(pages, start=1)
        if lines or index == 1
    ]


def read_pdf(path: Path, doc_id: str) -> list[Page]:
    reader = PdfReader(str(path))
    if not reader.pages:
        raise DocumentError(f"{path.name} 没有页面")
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(Page(doc_id, index, page.extract_text() or ""))
    return pages


def ocr_image(path: Path) -> str:
    executable = shutil.which("tesseract")
    if not executable:
        return ""
    completed = subprocess.run(
        [executable, str(path), "stdout", "-l", "chi_sim+eng", "--psm", "6"],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if completed.returncode != 0:
        return ""
    return completed.stdout or ""


def read_jpeg(path: Path, doc_id: str) -> list[Page]:
    return [Page(doc_id, 1, ocr_image(path))]


def read_pages(path: Path, doc_id: str) -> list[Page]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path, doc_id)
    if suffix == ".docx":
        return read_docx(path, doc_id)
    if suffix in {".jpg", ".jpeg"}:
        return read_jpeg(path, doc_id)
    raise DocumentError(f"暂不支持 {suffix or '无后缀'}：{path.name}。请用 PDF、docx 或 JPEG。")


def load_documents(paths: list[Path]) -> list[LoadedDocument]:
    loaded: list[LoadedDocument] = []
    for index, path in enumerate(paths, start=1):
        if not path.is_file():
            raise DocumentError(f"找不到文件：{path}")
        suffix = path.suffix.lower()
        mime = MIME.get(suffix)
        if mime is None:
            raise DocumentError(f"暂不支持 {suffix or '无后缀'}：{path.name}。请用 PDF、docx 或 JPEG。")
        doc_id = f"d{index}"
        document = LoadedDocument(
            doc_id=doc_id,
            path=path.resolve(),
            filename=path.name,
            sha256=file_sha256(path),
            mime=mime,
        )
        document.pages = read_pages(path, doc_id)
        if not document.pages:
            document.pages = [Page(doc_id, 1, "")]
        loaded.append(document)
    if not loaded:
        raise DocumentError("至少要有一份资料")
    return loaded
