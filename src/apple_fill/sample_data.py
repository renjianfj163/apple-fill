"""Write the local sample PDF, Word file, and JPEG used by the demo."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from PIL import Image, ImageDraw, ImageFont

PAGE_ONE = (
    "产品名称：青苹空气净化器",
    "品牌：青苹",
    "型号：QP-200",
    "制造商：青苹电器有限公司",
    "制造商地址：杭州市西湖区示例路 1 号",
    "联系人：林知夏",
    "联系电话：0571-88001000",
    "联系邮箱：sales@qingping.example",
    "额定电压：220V",
    "额定功率：35W",
    "执行标准：GB/T 18801-2015",
)
PAGE_TWO = "产品简介：用于室内空气过滤的家用电器"

_FONTS = (
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\msyh.ttf",
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simsun.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
)


def find_font() -> Path | None:
    for candidate in _FONTS:
        path = Path(candidate)
        if path.is_file():
            return path
    return None


def write_docx(path: Path) -> None:
    document = Document()
    for line in PAGE_ONE:
        document.add_paragraph(line)
    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph(PAGE_TWO)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def write_pdf(path: Path, font: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("body", fname=str(font))
    pdf.set_font("body", size=14)
    pdf.add_page()
    for line in PAGE_ONE:
        pdf.cell(0, 10, text=line, new_x="LMARGIN", new_y="NEXT")
    pdf.add_page()
    pdf.cell(0, 10, text=PAGE_TWO, new_x="LMARGIN", new_y="NEXT")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


def write_jpeg(path: Path, font: Path | None) -> None:
    image = Image.new("RGB", (1000, 1280), "white")
    if font is not None:
        draw = ImageDraw.Draw(image)
        face = ImageFont.truetype(str(font), 32)
        top = 48
        for line in PAGE_ONE:
            draw.text((48, top), line, fill="black", font=face)
            top += 52
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, "JPEG", quality=90)


def write_samples(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    docx_path = directory / "spec.docx"
    pdf_path = directory / "spec.pdf"
    image_path = directory / "label.jpg"
    write_docx(docx_path)
    font = find_font()
    paths = [docx_path]
    if font is not None:
        write_pdf(pdf_path, font)
        paths.append(pdf_path)
    write_jpeg(image_path, font)
    paths.append(image_path)
    return paths
